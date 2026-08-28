# -*- coding: utf-8 -*-
"""Reformattage intelligent d'un CV mal structure.

Quand un CV est trop mal mis en forme pour etre analyse automatiquement
(voir `docx_tools.assess_format`), on ne tente pas de l'ameliorer en place :
on reconstruit un CV propre a partir du texte brut, avec une rubrique
detectee intelligemment (meme sans intitule standard, via
`docx_tools.classify_heading_from_text`), et on laisse l'utilisateur choisir
un style visuel. L'orthographe/grammaire est corrigee par defaut au passage.
"""

import io
import re

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from . import corrector
from .docx_tools import classify_heading_from_text

# --------------------------------------------------------------------------- #
# Styles proposes a l'utilisateur.
# --------------------------------------------------------------------------- #
STYLES = [
    {
        "key": "classique",
        "label": "Classique",
        "description": "Titres en majuscules soulignés d'un filet fin. Sobre "
        "et professionnel, adapté à tous les secteurs.",
        "primary": "1F3864",
        "accent": "2E74B5",
        "uppercase_headings": True,
        "border_headings": True,
    },
    {
        "key": "moderne",
        "label": "Moderne",
        "description": "Titres colorés avec bandeau d'accent. Aéré, look "
        "actuel, adapté au digital, à la tech et au marketing.",
        "primary": "0F766E",
        "accent": "14B8A6",
        "uppercase_headings": True,
        "border_headings": True,
    },
    {
        "key": "minimaliste",
        "label": "Minimaliste",
        "description": "Noir et blanc, typographie fine, très épuré, sans "
        "fioriture.",
        "primary": "222222",
        "accent": "666666",
        "uppercase_headings": False,
        "border_headings": False,
    },
]
STYLE_MAP = {s["key"]: s for s in STYLES}

CANONICAL_TITLES = {
    "profil": "Profil",
    "experience": "Expérience professionnelle",
    "competences": "Compétences",
    "formation": "Formation",
    "langues": "Langues",
    "projets": "Projets",
    "contact": "Contact",
    "loisirs": "Centres d'intérêt",
    "certifications": "Certifications",
    "references": "Références",
}

_BULLET_RE = re.compile(r"^[•\-\*\u2013\u2022\u25AA\u25CF\u00B7‣]\s*")
_EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
_PHONE_RE = re.compile(r"(\+?\d[\d .\-()]{7,}\d)")


def _strip_bullet(line):
    return _BULLET_RE.sub("", line).strip()


def _looks_like_contact(line):
    low = line.lower()
    if _EMAIL_RE.search(line) or _PHONE_RE.search(line):
        return True
    if "linkedin" in low or "github" in low or "adresse" in low:
        return True
    if "|" in line and len(line) < 120:
        return True
    return False


def parse_cv(docx_bytes):
    """Analyse le texte brut d'un CV et le decoupe en rubriques.

    Retourne {"name": str, "contact_lines": [str], "sections": [...]}.
    """
    from .docx_tools import extract_raw_lines

    lines = [l for l in extract_raw_lines(docx_bytes) if l.strip()]
    if not lines:
        return {"name": "", "contact_lines": [], "sections": []}

    name = lines[0].strip()
    idx = 1
    contact_lines = []
    while idx < len(lines) and idx <= 5:
        line = lines[idx]
        is_h, _cat = classify_heading_from_text(line)
        if is_h:
            break
        if _looks_like_contact(line) or idx <= 2:
            contact_lines.append(line)
            idx += 1
        else:
            break

    sections = []
    current = None
    for line in lines[idx:]:
        is_h, cat = classify_heading_from_text(line)
        if is_h:
            current = {"title": line.strip(" :"), "category": cat or "autre", "items": []}
            sections.append(current)
            continue
        if current is None:
            current = {"title": "Informations", "category": "autre", "items": []}
            sections.append(current)
        current["items"].append(_strip_bullet(line))

    return {"name": name, "contact_lines": contact_lines, "sections": sections}


# --------------------------------------------------------------------------- #
# Construction du document reformatte.
# --------------------------------------------------------------------------- #
def _add_bottom_border(paragraph, color_hex, size=6):
    p = paragraph._p
    pPr = p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color_hex)
    pBdr.append(bottom)
    pPr.append(pBdr)


def build_docx(data, style_key="classique", correct=True):
    """Construit un .docx propre a partir des donnees parsees par `parse_cv`.

    Retourne (docx_bytes, nb_ameliorations_orthographe_grammaire).
    """
    style = STYLE_MAP.get(style_key, STYLES[0])
    primary = style["primary"]
    accent = style["accent"]
    total_corrections = 0

    def _fix(text):
        nonlocal total_corrections
        if not correct:
            return text
        fixed, n = corrector.correct_text(text)
        total_corrections += n
        return fixed

    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(1.6)
        section.bottom_margin = Cm(1.4)
        section.left_margin = Cm(1.8)
        section.right_margin = Cm(1.8)

    style_normal = doc.styles["Normal"]
    style_normal.font.name = "Calibri"
    style_normal.font.size = Pt(10.5)

    # ---- Nom ---- (nom propre : nettoyage minimal, pas de correction orthographique)
    name_text = corrector.light_clean(data.get("name", ""))
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(name_text.upper() if style_key != "minimaliste" else name_text)
    run.bold = True
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor.from_string(primary)

    # ---- Contact ---- (email/telephone/adresse : nettoyage minimal uniquement)
    if data.get("contact_lines"):
        contact_text = corrector.light_clean("  |  ".join(data["contact_lines"]))
        cp = doc.add_paragraph()
        cp.paragraph_format.space_after = Pt(10)
        crun = cp.add_run(contact_text)
        crun.font.size = Pt(10)
        crun.font.color.rgb = RGBColor.from_string("555555")

    # ---- Rubriques ----
    for sec in data.get("sections", []):
        cat = sec.get("category")
        title = CANONICAL_TITLES.get(cat) or _fix(sec["title"]).strip(" :").capitalize()

        hp = doc.add_paragraph()
        hp.paragraph_format.space_before = Pt(14)
        hp.paragraph_format.space_after = Pt(4)
        hrun = hp.add_run(title.upper() if style["uppercase_headings"] else title)
        hrun.bold = True
        hrun.font.size = Pt(12.5)
        hrun.font.color.rgb = RGBColor.from_string(primary)
        if style["border_headings"]:
            _add_bottom_border(hp, accent)

        for item in sec.get("items", []):
            text = _fix(item)
            if not text:
                continue
            ip = doc.add_paragraph(style="List Bullet")
            ip.paragraph_format.space_after = Pt(2)
            irun = ip.add_run(text)
            irun.font.size = Pt(10.5)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue(), total_corrections
