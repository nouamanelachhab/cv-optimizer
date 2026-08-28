"""docgen: docxtpl-based document regeneration.

This module replaces "edit paragraph runs in place" with the pipeline
mandated by the addendum: ``docx -> typed model -> (enrichment) -> docx``.

Step 1 (this delivery): build a Jinja2 template `.docx` from an ingested
``Document`` (one tag per block), then re-render it with the *original*
block texts as context, so that regenerating a CV from its own model
reproduces the source content exactly. No optimization/enrichment logic
lives here -- this is the baseline that later steps will build on.

``jinja2.StrictUndefined`` is used so that a missing context value raises
``UndefinedError`` at render time instead of silently leaving an unfilled
placeholder in the output (this is what makes defect D1 -- unresolved
placeholders like "(e) par responsable" -- structurally impossible).
"""

from __future__ import annotations

import io

from docx import Document as DocxDocument
from docxtpl import DocxTemplate
from jinja2 import Environment, StrictUndefined

from src.ingest.models import Bullet, Document

__all__ = ["tag_name", "build_template", "baseline_context", "render"]

_FALLBACK_BULLET_STYLE = "List Bullet"
_FALLBACK_HEADING_STYLE = "Heading 1"


def tag_name(block_id: str) -> str:
    """Jinja2/docxtpl context key for a given block id.

    Prefixed with ``block_`` so the result is always a valid Jinja
    identifier, even though ``block_id`` (a truncated sha1 hex digest) may
    start with a digit.
    """
    return f"block_{block_id}"


def _add_paragraph(docx_document: DocxDocument, block: Bullet, fallback_style: str | None) -> None:
    text = "{{ " + tag_name(block.block_id) + " }}"
    style_name = block.style_name or fallback_style
    if style_name:
        try:
            docx_document.add_paragraph(text, style=style_name)
            return
        except KeyError:
            pass  # style not present in this template's style catalogue
    docx_document.add_paragraph(text)


def build_template(doc: Document) -> bytes:
    """Build a docxtpl-compatible `.docx` template with one Jinja tag per block.

    Each paragraph carries exactly one tag in a single run, so docxtpl can
    always resolve it (split runs are the classic cause of unresolved tags
    in Jinja-based docx templating).
    """
    docx_document = DocxDocument()

    if doc.name_block is not None:
        _add_paragraph(docx_document, doc.name_block, None)
    for contact_block in doc.contact_blocks:
        _add_paragraph(docx_document, contact_block, None)

    for section in doc.sections:
        heading_block = Bullet(
            block_id=section.block_id,
            text=section.heading_text,
            xml_offset=section.xml_offset,
        )
        _add_paragraph(docx_document, heading_block, _FALLBACK_HEADING_STYLE)

        for entry in section.entries:
            if entry.text:
                entry_block = Bullet(
                    block_id=entry.block_id,
                    text=entry.text,
                    xml_offset=entry.xml_offset,
                    style_name=entry.style_name,
                )
                _add_paragraph(docx_document, entry_block, None)
            for bullet in entry.bullets:
                _add_paragraph(docx_document, bullet, _FALLBACK_BULLET_STYLE)

        for free_text_block in section.free_text_blocks:
            _add_paragraph(docx_document, free_text_block, None)

    buffer = io.BytesIO()
    docx_document.save(buffer)
    return buffer.getvalue()


def baseline_context(doc: Document) -> dict[str, str]:
    """Context mapping every block to its own original text (identity regeneration)."""
    return {tag_name(block.block_id): block.text for block in doc.all_blocks()}


def render(template_bytes: bytes, context: dict[str, str]) -> bytes:
    """Render a docxtpl template with ``context``, failing loudly on missing keys.

    Uses ``StrictUndefined``: any tag not present in ``context`` raises
    ``jinja2.exceptions.UndefinedError`` instead of producing an unresolved
    or blank placeholder in the output document.
    """
    template = DocxTemplate(io.BytesIO(template_bytes))
    env = Environment(undefined=StrictUndefined)
    template.render(context, jinja_env=env)
    buffer = io.BytesIO()
    template.save(buffer)
    return buffer.getvalue()
