# -*- coding: utf-8 -*-
"""Fixtures partagees pour les tests du moteur ATS deterministe."""

import io

import pytest
from docx import Document as DocxDocument


def _make_docx(paragraphs_spec):
    """Construit un .docx en memoire a partir d'une liste de
    (texte, style_ou_None) et retourne les octets."""
    doc = DocxDocument()
    for text, style in paragraphs_spec:
        if style:
            doc.add_paragraph(text, style=style)
        else:
            doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture
def simple_cv_bytes():
    spec = [
        ("Jean Dupont", None),
        ("jean.dupont@email.com | 06 12 34 56 78 | Paris", None),
        ("PROFIL", None),
        ("Développeur backend avec 5 ans d'expérience en Python et SQL.", None),
        ("EXPERIENCE", None),
        ("Développeur Python — Decathlon — 2020-2023", None),
        ("Développement et maintenance d'une application de gestion des stocks.", "List Bullet"),
        ("Optimisation des traitements Oracle PL/SQL (-65% de temps).", "List Bullet"),
        ("COMPETENCES", None),
        ("Python, SQL, Docker, Anglais", None),
        ("FORMATION", None),
        ("Ingénieur diplômé en génie informatique — 2019", None),
    ]
    return _make_docx(spec)


@pytest.fixture
def offer_text_decathlon():
    return (
        "Decathlon recherche un Développeur Python (H/F).\n\n"
        "Missions : développement d'applications de gestion des stocks, "
        "optimisation de requêtes Oracle PL/SQL, exploitation d'un infocentre, "
        "animation d'ateliers avec les équipes.\n\n"
        "Profil recherché : compétences indispensables en Python, SQL, Docker. "
        "La maîtrise de Kubernetes est un plus. Anglais courant souhaité.\n\n"
        "Compétences : Python, SQL, Oracle, PL/SQL, Docker, Kubernetes, ISO 27001."
    )
