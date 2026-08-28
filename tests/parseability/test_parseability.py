"""Tests for the parseability module."""

from __future__ import annotations

import io

from docx import Document as DocxDocument

from src.ingest import parse_docx
from src.parseability import run_parseability_audit
from src.parseability.checks import check_double_parser_diff


def test_out_of_flow_content_flagged_when_present(simple_cv_bytes):
    doc = parse_docx(simple_cv_bytes, source_name="cv.docx")
    report = run_parseability_audit(doc)
    codes = {i.code for i in report.issues}
    if doc.out_of_flow_blocks:
        assert "out_of_flow_content" in codes


def test_unrecognized_heading_flagged():
    from tests.conftest import _make_docx

    docx_bytes = _make_docx(
        [
            ("Jean Dupont", None),
            ("RUBRIQUE BIZARRE", "Heading 1"),
            ("Un paragraphe quelconque.", None),
        ]
    )
    doc = parse_docx(docx_bytes, source_name="cv.docx")
    report = run_parseability_audit(doc)
    codes = {i.code for i in report.issues}
    assert "unrecognized_heading" in codes


def test_recognized_headings_not_flagged(simple_cv_bytes):
    doc = parse_docx(simple_cv_bytes, source_name="cv.docx")
    report = run_parseability_audit(doc)
    codes = {i.code for i in report.issues}
    assert "unrecognized_heading" not in codes


def test_inconsistent_date_format_detected():
    from tests.conftest import _make_docx

    docx_bytes = _make_docx(
        [
            ("Jean Dupont", None),
            ("EXPERIENCE", "Heading 1"),
            ("Ingénieur - 01/2020 à 06/2021", None),
            ("Ingénieur - janvier 2018 à décembre 2019", None),
        ]
    )
    doc = parse_docx(docx_bytes, source_name="cv.docx")
    report = run_parseability_audit(doc)
    codes = {i.code for i in report.issues}
    assert "inconsistent_date_format" in codes


def test_non_chronological_order_detected():
    from tests.conftest import _make_docx

    docx_bytes = _make_docx(
        [
            ("Jean Dupont", None),
            ("EXPERIENCE", "Heading 1"),
            ("Ingénieur - 2018 à 2019", None),
            ("Ingénieur - 2020 à 2021", None),
        ]
    )
    doc = parse_docx(docx_bytes, source_name="cv.docx")
    report = run_parseability_audit(doc)
    codes = {i.code for i in report.issues}
    assert "non_chronological_order" in codes


def test_docx_bytes_checks_skipped_when_not_provided(simple_cv_bytes):
    doc = parse_docx(simple_cv_bytes, source_name="cv.docx")
    report = run_parseability_audit(doc, docx_bytes=None)
    codes = {i.code for i in report.issues}
    assert "multi_column_layout" not in codes
    assert "hyperlink_mismatch" not in codes


def test_multi_column_and_hyperlink_checks_run_when_bytes_provided(simple_cv_bytes):
    doc = parse_docx(simple_cv_bytes, source_name="cv.docx")
    # Should not raise even though the fixture has no columns/hyperlinks.
    report = run_parseability_audit(doc, docx_bytes=simple_cv_bytes)
    assert isinstance(report.issues, list)


def _make_docx_with_header_only_content(header_text: str) -> bytes:
    docx_document = DocxDocument()
    docx_document.add_paragraph("Jean Dupont")
    docx_document.add_paragraph("Un CV sans contact dans le corps du document.")
    section = docx_document.sections[0]
    section.header.is_linked_to_previous = False
    section.header.paragraphs[0].text = header_text
    buffer = io.BytesIO()
    docx_document.save(buffer)
    return buffer.getvalue()


def test_double_parser_diff_detects_content_only_in_header():
    docx_bytes = _make_docx_with_header_only_content("Contact exclusif en en-tête uniquement")
    issues = check_double_parser_diff(docx_bytes)
    codes = {i.code for i in issues}
    assert "double_parser_content_loss" in codes


def test_double_parser_diff_escalates_email_lost_in_header():
    docx_bytes = _make_docx_with_header_only_content("jean.dupont@email.com")
    issues = check_double_parser_diff(docx_bytes)
    codes = {i.code for i in issues}
    assert "double_parser_critical_content_loss" in codes


def test_double_parser_diff_empty_when_body_contains_everything(simple_cv_bytes):
    issues = check_double_parser_diff(simple_cv_bytes)
    codes = {i.code for i in issues}
    assert "double_parser_critical_content_loss" not in codes


def test_double_parser_diff_wired_into_full_audit():
    docx_bytes = _make_docx_with_header_only_content("06 12 34 56 78")
    doc = parse_docx(docx_bytes, source_name="cv.docx")
    report = run_parseability_audit(doc, docx_bytes=docx_bytes)
    codes = {i.code for i in report.issues}
    assert "double_parser_critical_content_loss" in codes
