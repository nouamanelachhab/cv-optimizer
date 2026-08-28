# -*- coding: utf-8 -*-
"""Tests du mécanisme d'édition par spans (src/docx_tools/spans.py).

Le test de round-trip (spans vides) est celui exigé en premier par le
prompt : tant qu'il échoue, aucune modification par span n'est sûre.
"""

from __future__ import annotations

import io

import pytest
from docx import Document as DocxDocument

from src.docx_tools.spans import Span, apply_spans, paragraph_text


def _build_docx(paragraphs_runs: list[list[tuple[str, dict]]]) -> bytes:
    """Construit un .docx où chaque paragraphe a des runs explicites, avec
    mise en forme distincte par run (pour vérifier qu'elle est préservée)."""
    doc = DocxDocument()
    for runs_spec in paragraphs_runs:
        p = doc.add_paragraph()
        for text, fmt in runs_spec:
            run = p.add_run(text)
            if fmt.get("bold"):
                run.bold = True
            if fmt.get("italic"):
                run.italic = True
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _load(docx_bytes: bytes) -> DocxDocument:
    return DocxDocument(io.BytesIO(docx_bytes))


def test_round_trip_empty_spans_produce_byte_identical_document():
    """Test à écrire en premier : appliquer une liste de spans vide sur
    chaque paragraphe d'un CV ne doit rien changer, octet pour octet."""
    original_bytes = _build_docx(
        [
            [("Expérience en ", {}), ("Python", {"bold": True}), (" et SQL.", {})],
            [("Ingénieur ", {"italic": True}), ("diplômé", {})],
        ]
    )
    doc = _load(original_bytes)
    for paragraph in doc.paragraphs:
        apply_spans(paragraph, [])

    buf = io.BytesIO()
    doc.save(buf)
    assert buf.getvalue() == original_bytes


def test_round_trip_noop_span_preserves_text_and_formatting():
    """Une span qui remplace un segment par lui-même ne doit rien changer au
    texte ni au run affecté (rPr intact)."""
    original_bytes = _build_docx(
        [[("Expérience en ", {}), ("Python", {"bold": True}), (" et SQL.", {})]]
    )
    doc = _load(original_bytes)
    paragraph = doc.paragraphs[0]
    text = paragraph_text(paragraph)
    apply_spans(paragraph, [Span(start=0, end=len(text), replacement=text)])

    assert paragraph_text(paragraph) == text
    assert paragraph.runs[1].bold is True


def test_apply_spans_rewrites_only_touched_run():
    """Remplacer un mot au milieu d'un run laisse les autres runs (et leur
    mise en forme) totalement intacts."""
    doc_bytes = _build_docx(
        [[("Expérience en ", {}), ("Java", {"bold": True}), (" et SQL.", {})]]
    )
    doc = _load(doc_bytes)
    paragraph = doc.paragraphs[0]
    text = paragraph_text(paragraph)
    start = text.index("Java")
    apply_spans(paragraph, [Span(start=start, end=start + len("Java"), replacement="Python")])

    assert paragraph_text(paragraph) == "Expérience en Python et SQL."
    assert paragraph.runs[0].text == "Expérience en "
    assert paragraph.runs[1].text == "Python"
    assert paragraph.runs[1].bold is True
    assert paragraph.runs[2].text == " et SQL."


def test_apply_spans_never_deletes_or_recreates_runs():
    """Le nombre de runs ne doit jamais changer, même quand une span
    remplace un segment entier d'un run par une chaîne vide ou plus longue."""
    doc_bytes = _build_docx(
        [[("A ", {}), ("B", {"bold": True}), (" C", {})]]
    )
    doc = _load(doc_bytes)
    paragraph = doc.paragraphs[0]
    n_runs_before = len(paragraph.runs)
    text = paragraph_text(paragraph)
    start = text.index("B")
    apply_spans(paragraph, [Span(start=start, end=start + 1, replacement="")])

    assert len(paragraph.runs) == n_runs_before
    assert paragraph_text(paragraph) == "A  C"


def test_apply_spans_spanning_multiple_runs_emits_replacement_once():
    """Une span qui chevauche plusieurs runs place le texte de remplacement
    une seule fois (dans le premier run touché) et vide les runs suivants
    pour la portion couverte."""
    doc_bytes = _build_docx(
        [[("Dévelop", {}), ("pement", {"bold": True}), (" web", {})]]
    )
    doc = _load(doc_bytes)
    paragraph = doc.paragraphs[0]
    text = paragraph_text(paragraph)
    start = text.index("Dévelop")
    end = start + len("Développement")
    apply_spans(paragraph, [Span(start=start, end=end, replacement="Conception")])

    assert paragraph_text(paragraph) == "Conception web"


def test_apply_spans_rejects_overlapping_spans():
    doc_bytes = _build_docx([[("abcdef", {})]])
    doc = _load(doc_bytes)
    paragraph = doc.paragraphs[0]
    with pytest.raises(ValueError):
        apply_spans(
            paragraph,
            [Span(start=0, end=3, replacement="x"), Span(start=2, end=5, replacement="y")],
        )
